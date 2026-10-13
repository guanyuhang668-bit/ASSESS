#!/root/UML图工具/uml-grader/backend/.venv/bin/python3
"""基底选择脚本：用预实验数据对比各 LLM 模型与教师评分的一致性。

用法:
    cd /root/UML图工具/uml-grader
    python scripts/benchmark_models.py

输出:
    scripts/benchmark_results/  下生成：
      - detail.csv            每条指标的 教师分 vs 各模型分
      - summary.csv           按模型 × 图类型汇总的 MAE / RMSE / Pearson / 一致率
      - report.txt            文字报告
"""

from __future__ import annotations

import asyncio
import csv
import json
import logging
import math
import os
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

# ---- 把 backend 加入 sys.path ----
BACKEND_DIR = Path(__file__).resolve().parent.parent / "backend"
sys.path.insert(0, str(BACKEND_DIR))

os.environ.setdefault("APP_ENV", "dev")

# 需要在 import app 模块之前设好 path
from app.services.grading.cot_runner import CotRunner
from app.services.grading.rubric_loader import load_rubric_from_yaml, RubricDef, CriterionDef
from app.services.llm.factory import load_all_clients
from app.services.parser_svc import parse_submission, pick_diagram
from app.uml_grader.grader.schema import CriterionScore

import openpyxl

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("benchmark")

# ======================= 配置 =======================

MDJ_ROOT = Path("/tmp/预实验_all")
GT_DIR = Path("/root/UML图工具/基底选择")
RUBRIC_DIR = BACKEND_DIR / "configs" / "rubrics"
OUTPUT_DIR = Path(__file__).resolve().parent / "benchmark_results"

DIAGRAM_TYPES = ["use_case", "class", "sequence", "swimlane"]
DIAGRAM_CN = {"use_case": "用例图", "class": "类图", "sequence": "时序图", "swimlane": "泳道图"}
DIAGRAM_FOLDER_KW = {"use_case": "用例图", "class": "类图", "sequence": "时序图", "swimlane": "泳道图"}

# XLSX 文件与学生名映射
GT_FILES = {
    "A": "UML指标A组.xlsx",
    "B": "UML指标B组.xlsx",
    "C": "UML指标c组 (1).xlsx",
    "D": "D组-uml指标.xlsx",
}

# XLSX 中学生名 → MDJ 目录名的映射（处理不一致）
STUDENT_NAME_MAP = {
    "A-黄彬菊": "A组-黄彬菊",
    "A-江羽萱": "A组-江羽萱",
    "A-唐雨恬": "A组-唐雨恬",
    "A-王天力": "A组-王天立",  # 力→立
    "A-郑伟": "A组-郑伟",
    "B-王俊晧": "B组-王俊晧",
    "B-唐亮婕": "B组-唐亮婕",
    "B-黄韬喆": None,  # 无 mdj
    "B-朱子鸣": "B组-朱子鸣",
    "B-代宇扬": None,  # 无 mdj
    "C 刘思雨": "C组-刘思语",  # 雨→语
    "C-刘思雨": "C组-刘思语",
    "C-袁瑛凯": "C组-袁瑛凯",
    "C-熊宇辉": "C组-熊宇辉",
    "C-冯梦雅": "C组-冯亚梦",  # 梦雅→亚梦
    "C-谭艺丰": "C组-谭艺丰",
    "D-黄汪洋": "D组-黄汪洋",
    "D-刘玉玲": "D组-刘玉玲",
    "D-史永康": "D组-史永康",
    "D-李昕妍": "D组-李昕妍 ",  # 注意末尾空格
    "D-卫欣": "D组-卫欣",
}


# ======================= 数据结构 =======================

@dataclass
class TeacherScore:
    student_name: str
    diagram_type: str
    code: str
    score: float | None  # None = not_applicable (×)


@dataclass
class ModelScore:
    student_name: str
    diagram_type: str
    model_name: str
    code: str
    score: float
    not_applicable: bool


# ======================= 读取教师评分 =======================

def _parse_diagram_section(ws, start_row: int, end_row: int, diagram_type: str,
                           student_cols: list[tuple[str, int]]) -> list[TeacherScore]:
    """解析一个图类型区域的教师评分。"""
    scores = []
    for row in range(start_row, end_row + 1):
        cell_b = ws.cell(row, 2).value
        if not cell_b:
            continue
        m = re.match(r"([A-Z]{2}\d+\.\d+)", str(cell_b))
        if not m:
            continue
        code = m.group(1)
        for stu_name, col in student_cols:
            val = ws.cell(row, col).value
            if val == "×" or val == "x" or val == "X":
                scores.append(TeacherScore(stu_name, diagram_type, code, None))
            elif isinstance(val, (int, float)):
                scores.append(TeacherScore(stu_name, diagram_type, code, float(val)))
            elif val is None:
                # 可能继承了上级的 ×，也视为 N/A
                scores.append(TeacherScore(stu_name, diagram_type, code, None))
            else:
                # 其他意外值
                logger.warning("Unexpected score value at row %d col %d: %r", row, col, val)
                scores.append(TeacherScore(stu_name, diagram_type, code, None))
    return scores


def load_teacher_scores() -> list[TeacherScore]:
    """从所有 XLSX 加载教师评分。"""
    all_scores: list[TeacherScore] = []

    for group, fname in GT_FILES.items():
        path = GT_DIR / fname
        wb = openpyxl.load_workbook(str(path), data_only=True)
        ws = wb["Sheet1"]

        # 找学生列
        student_cols: list[tuple[str, int]] = []
        for c in range(1, ws.max_column + 1):
            if ws.cell(2, c).value == "分数":
                for cc in range(c, 0, -1):
                    if ws.cell(1, cc).value:
                        student_cols.append((str(ws.cell(1, cc).value).strip(), c))
                        break

        # 找各图类型的行范围
        sections: list[tuple[str, int, int]] = []  # (diagram_type, start_row, end_row)
        diagram_rows: list[tuple[str, int]] = []
        for i in range(1, ws.max_row + 1):
            val_a = ws.cell(i, 1).value
            if val_a and str(val_a).strip() in DIAGRAM_CN.values():
                cn = str(val_a).strip()
                dt = [k for k, v in DIAGRAM_CN.items() if v == cn][0]
                diagram_rows.append((dt, i))

        for idx, (dt, start) in enumerate(diagram_rows):
            if idx + 1 < len(diagram_rows):
                end = diagram_rows[idx + 1][1] - 1
            else:
                end = ws.max_row
            # 实际数据从 start+2 开始（跳过标题行）
            sections.append((dt, start + 2, end))

        for dt, sr, er in sections:
            scores = _parse_diagram_section(ws, sr, er, dt, student_cols)
            all_scores.extend(scores)

    logger.info("Loaded %d teacher scores", len(all_scores))
    return all_scores


# ======================= 查找学生 MDJ 文件 =======================

_MDJ_DTYPE_MAP = {
    "UMLUseCaseDiagram": "use_case",
    "UMLClassDiagram": "class",
    "UMLSequenceDiagram": "sequence",
    "UMLActivityDiagram": "swimlane",
}


def _detect_diagram_type(mdj_path: Path) -> str | None:
    """通过解析 MDJ 内容检测图类型。"""
    try:
        with open(mdj_path, encoding="utf-8") as f:
            data = json.load(f)
        diagrams = parse_submission(data)
        if diagrams:
            return diagrams[0].diagram_type
    except Exception:
        pass
    return None


def find_mdj(student_folder_name: str, diagram_type: str) -> Path | None:
    """在 MDJ_ROOT/<student_folder_name>/ 下递归查找对应图类型的 mdj 文件。"""
    base = MDJ_ROOT / student_folder_name
    if not base.exists():
        return None
    keyword = DIAGRAM_FOLDER_KW[diagram_type]
    # 优先：路径中含关键词
    for p in base.rglob("*.mdj"):
        if keyword in str(p):
            return p
    # fallback：解析 MDJ 内容匹配图类型
    for p in base.rglob("*.mdj"):
        detected = _detect_diagram_type(p)
        if detected == diagram_type:
            return p
    return None


# ======================= 运行评分 =======================

async def grade_one(
    client, rubric: RubricDef, mdj_path: Path, diagram_type: str,
    student_name: str,
) -> list[ModelScore]:
    """用一个 LLM 对一个学生的一种图评分。"""
    with open(mdj_path, encoding="utf-8") as f:
        raw_data = json.load(f)

    diagrams = parse_submission(raw_data)
    diagram = pick_diagram(diagrams, diagram_type)
    if diagram is None:
        logger.warning("[%s] %s: no %s diagram found in %s", client.name, student_name, diagram_type, mdj_path)
        return []

    runner = CotRunner(
        client=client,
        rubric=rubric,
        diagram=diagram,
        problem_statement="",  # 预实验无题目文本
        student_id=student_name,
    )

    try:
        result = await runner.run()
    except Exception as e:
        logger.error("[%s] %s/%s failed: %s", client.name, student_name, diagram_type, e)
        return []

    scores = []
    for cs in result.criteria:
        scores.append(ModelScore(
            student_name=student_name,
            diagram_type=diagram_type,
            model_name=client.name,
            code=cs.code,
            score=cs.score,
            not_applicable=cs.not_applicable,
        ))
    return scores


# ======================= 统计指标 =======================

def compute_metrics(teacher: list[float], model: list[float]) -> dict[str, float]:
    """计算 MAE, RMSE, Pearson 相关系数, 一致率（误差≤1分）。"""
    n = len(teacher)
    if n == 0:
        return {"mae": 0, "rmse": 0, "pearson": 0, "agree_1": 0, "agree_2": 0, "n": 0}

    diffs = [abs(t - m) for t, m in zip(teacher, model)]
    mae = sum(diffs) / n
    rmse = math.sqrt(sum(d ** 2 for d in diffs) / n)

    # 一致率
    agree_1 = sum(1 for d in diffs if d <= 1) / n
    agree_2 = sum(1 for d in diffs if d <= 2) / n

    # Pearson
    mean_t = sum(teacher) / n
    mean_m = sum(model) / n
    cov = sum((t - mean_t) * (m - mean_m) for t, m in zip(teacher, model))
    std_t = math.sqrt(sum((t - mean_t) ** 2 for t in teacher))
    std_m = math.sqrt(sum((m - mean_m) ** 2 for m in model))
    pearson = cov / (std_t * std_m) if std_t > 0 and std_m > 0 else 0.0

    return {"mae": mae, "rmse": rmse, "pearson": pearson, "agree_1": agree_1, "agree_2": agree_2, "n": n}


# ======================= 增量持久化 =======================

DETAIL_FIELDS = ["student", "diagram_type", "code", "teacher_score", "model", "model_score", "model_na"]


def _load_done_keys() -> set[tuple[str, str, str]]:
    """从已有 detail.csv 读取已完成的 (student, diagram_type, model) 组合。"""
    detail_path = OUTPUT_DIR / "detail.csv"
    done: set[tuple[str, str, str]] = set()
    if not detail_path.exists():
        return done
    with open(detail_path, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            done.add((row["student"], row["diagram_type"], row["model"]))
    return done


def _ensure_detail_header():
    """确保 detail.csv 存在且有表头。"""
    detail_path = OUTPUT_DIR / "detail.csv"
    if detail_path.exists() and detail_path.stat().st_size > 0:
        return
    with open(detail_path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerow(DETAIL_FIELDS)


def _append_detail_rows(rows: list[dict]):
    """追加行到 detail.csv。"""
    detail_path = OUTPUT_DIR / "detail.csv"
    with open(detail_path, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=DETAIL_FIELDS)
        w.writerows(rows)


def _load_all_detail() -> list[dict]:
    """加载完整的 detail.csv。"""
    detail_path = OUTPUT_DIR / "detail.csv"
    if not detail_path.exists():
        return []
    with open(detail_path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def generate_reports(teacher_map: dict[tuple[str, str, str], float | None]):
    """从 detail.csv 重新生成 summary.csv 和 report.txt。"""
    all_rows = _load_all_detail()
    if not all_rows:
        logger.warning("detail.csv is empty, no report to generate")
        return

    model_names = sorted(set(r["model"] for r in all_rows))

    summary_rows = []
    for model_name in model_names:
        for dt in DIAGRAM_TYPES:
            teacher_vals, model_vals = [], []
            na_match = na_total = 0
            for row in all_rows:
                if row["model"] != model_name or row["diagram_type"] != dt:
                    continue
                ts_val = teacher_map.get((row["student"], row["diagram_type"], row["code"]))
                m_score = float(row["model_score"]) if row["model_score"] else 0
                m_na = row["model_na"] == "True"
                if ts_val is None:
                    na_total += 1
                    if m_na:
                        na_match += 1
                    continue
                if m_na:
                    model_vals.append(0.0)
                    teacher_vals.append(ts_val)
                else:
                    teacher_vals.append(ts_val)
                    model_vals.append(m_score)

            metrics = compute_metrics(teacher_vals, model_vals)
            na_acc = na_match / na_total if na_total > 0 else float("nan")
            summary_rows.append({
                "model": model_name, "diagram_type": dt, "n": metrics["n"],
                "mae": round(metrics["mae"], 3), "rmse": round(metrics["rmse"], 3),
                "pearson": round(metrics["pearson"], 3),
                "agree_1pt": round(metrics["agree_1"] * 100, 1),
                "agree_2pt": round(metrics["agree_2"] * 100, 1),
                "na_total": na_total, "na_match": na_match,
                "na_accuracy": round(na_acc * 100, 1) if not math.isnan(na_acc) else "N/A",
            })

        # 整体
        all_t, all_m = [], []
        for row in all_rows:
            if row["model"] != model_name:
                continue
            ts_val = teacher_map.get((row["student"], row["diagram_type"], row["code"]))
            m_na = row["model_na"] == "True"
            if ts_val is None or m_na:
                continue
            all_t.append(ts_val)
            all_m.append(float(row["model_score"]))
        overall = compute_metrics(all_t, all_m)
        summary_rows.append({
            "model": model_name, "diagram_type": "ALL", "n": overall["n"],
            "mae": round(overall["mae"], 3), "rmse": round(overall["rmse"], 3),
            "pearson": round(overall["pearson"], 3),
            "agree_1pt": round(overall["agree_1"] * 100, 1),
            "agree_2pt": round(overall["agree_2"] * 100, 1),
            "na_total": "", "na_match": "", "na_accuracy": "",
        })

    # 写 summary.csv
    summary_path = OUTPUT_DIR / "summary.csv"
    with open(summary_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=[
            "model", "diagram_type", "n", "mae", "rmse", "pearson",
            "agree_1pt", "agree_2pt", "na_total", "na_match", "na_accuracy",
        ])
        w.writeheader()
        w.writerows(summary_rows)
    logger.info("Summary CSV: %s", summary_path)

    # 文字报告
    report_lines = ["=" * 60, "模型基底选择评估报告", "=" * 60, ""]
    report_lines.append("各模型整体表现（按 MAE 排序，越低越好）：")
    report_lines.append("-" * 60)
    report_lines.append(f"{'模型':<20} {'MAE':>6} {'RMSE':>6} {'Pearson':>8} {'±1分准确率':>10} {'±2分准确率':>10} {'样本数':>6}")
    report_lines.append("-" * 60)

    overall_rows = sorted(
        [r for r in summary_rows if r["diagram_type"] == "ALL"],
        key=lambda r: r["mae"],
    )
    for r in overall_rows:
        report_lines.append(
            f"{r['model']:<20} {r['mae']:>6.3f} {r['rmse']:>6.3f} {r['pearson']:>8.3f} "
            f"{r['agree_1pt']:>9.1f}% {r['agree_2pt']:>9.1f}% {r['n']:>6}"
        )

    report_lines.append("")
    report_lines.append("=" * 60)
    report_lines.append("按图类型分解：")
    for dt in DIAGRAM_TYPES:
        report_lines.append(f"\n--- {DIAGRAM_CN[dt]} ---")
        dt_rows = sorted(
            [r for r in summary_rows if r["diagram_type"] == dt],
            key=lambda r: r["mae"],
        )
        for r in dt_rows:
            report_lines.append(
                f"  {r['model']:<20} MAE={r['mae']:.3f}  RMSE={r['rmse']:.3f}  "
                f"Pearson={r['pearson']:.3f}  ±1分={r['agree_1pt']:.1f}%  "
                f"±2分={r['agree_2pt']:.1f}%  N/A识别={r['na_accuracy']}"
            )

    report_lines.append("")
    report_lines.append("=" * 60)
    best = overall_rows[0] if overall_rows else None
    if best:
        report_lines.append(f"★ 推荐模型: {best['model']}")
        report_lines.append(f"  理由: MAE 最低 ({best['mae']:.3f})，与教师评分最接近")

    report_text = "\n".join(report_lines)
    report_path = OUTPUT_DIR / "report.txt"
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report_text)
    print("\n" + report_text)
    logger.info("Report: %s", report_path)


# ======================= 主流程 =======================

async def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1. 加载教师评分
    teacher_scores = load_teacher_scores()
    teacher_map: dict[tuple[str, str, str], float | None] = {}
    for ts in teacher_scores:
        teacher_map[(ts.student_name, ts.diagram_type, ts.code)] = ts.score

    # 2. 加载评分标准
    rubrics: dict[str, RubricDef] = {}
    for dt in DIAGRAM_TYPES:
        rubrics[dt] = load_rubric_from_yaml(RUBRIC_DIR / f"{dt}.yaml")

    # 3. 加载 LLM 客户端
    clients = load_all_clients(BACKEND_DIR / "configs" / "models.yaml")
    clients = {k: v for k, v in clients.items() if k != "mock"}
    logger.info("Models to benchmark: %s", list(clients.keys()))

    if not clients:
        logger.error("No LLM clients available (check API keys in .env)")
        return

    # 4. 收集所有 (student, diagram_type) 组合
    tasks: list[tuple[str, str, str, Path]] = []
    student_names_seen = set()
    for ts in teacher_scores:
        key = (ts.student_name, ts.diagram_type)
        if key in student_names_seen:
            continue
        student_names_seen.add(key)

        folder = STUDENT_NAME_MAP.get(ts.student_name)
        if folder is None:
            logger.info("Skip %s (no mdj files)", ts.student_name)
            continue

        mdj = find_mdj(folder, ts.diagram_type)
        if mdj is None:
            logger.warning("Skip %s/%s: mdj not found in %s", ts.student_name, ts.diagram_type, folder)
            continue
        tasks.append((ts.student_name, folder, ts.diagram_type, mdj))

    # 5. 加载已完成进度（断点续跑）
    _ensure_detail_header()
    done_keys = _load_done_keys()
    total = len(tasks) * len(clients)
    skipped = sum(1 for m in clients for s, _, d, _ in tasks if (s, d, m) in done_keys)
    logger.info("Total tasks: %d, already done: %d, remaining: %d", total, skipped, total - skipped)

    # 6. 执行评分（增量保存，每完成一个任务立即写入 detail.csv）
    for model_name, client in clients.items():
        logger.info("=== Benchmarking model: %s ===", model_name)
        for i, (stu_name, folder, dt, mdj_path) in enumerate(tasks):
            if (stu_name, dt, model_name) in done_keys:
                logger.info("  [%d/%d] %s / %s — SKIP (already done)", i + 1, len(tasks), stu_name, dt)
                continue

            logger.info("  [%d/%d] %s / %s", i + 1, len(tasks), stu_name, dt)
            t0 = time.time()
            scores = await grade_one(client, rubrics[dt], mdj_path, dt, stu_name)
            elapsed = time.time() - t0
            logger.info("    → %d criteria scored in %.1fs", len(scores), elapsed)

            if scores:
                rows = []
                for ms in scores:
                    ts_val = teacher_map.get((ms.student_name, ms.diagram_type, ms.code))
                    rows.append({
                        "student": ms.student_name,
                        "diagram_type": ms.diagram_type,
                        "code": ms.code,
                        "teacher_score": "" if ts_val is None else ts_val,
                        "model": ms.model_name,
                        "model_score": ms.score,
                        "model_na": ms.not_applicable,
                    })
                _append_detail_rows(rows)
                done_keys.add((stu_name, dt, model_name))

    # 7. 生成报告
    logger.info("All tasks done. Generating reports...")
    generate_reports(teacher_map)


if __name__ == "__main__":
    asyncio.run(main())
