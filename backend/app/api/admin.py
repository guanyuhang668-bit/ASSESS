"""管理员接口：批量导入学生账号等。"""

from __future__ import annotations

import csv
import json
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from ..core.config import get_settings
from ..core.db import get_session
from ..core.minio_client import get_minio_client, get_minio_url
from ..core.security import hash_password, require_role
from ..models.chat import ChatMessage, ChatSession, ChatCriterionScore
from ..models.chat_teacher_score import ChatTeacherIssue, ChatTeacherScore
from ..models.class_ import SchoolClass
from ..models.grading import CriterionScore, GradingResult
from ..models.student import Student
from ..models.submission import Submission
from ..models.teacher_score import TeacherScore
from ..models.user import AppUser
from ..schemas.chat_teacher import (
    ChatTeacherIssueOut,
    ChatTeacherScoreIn,
    ChatTeacherScoreOut,
)

router = APIRouter()

require_admin = require_role("admin")


def _safe_path_part(value: object, fallback: str = "unknown") -> str:
    text = str(value or fallback).strip()
    text = re.sub(r'[\\/:*?"<>|\s]+', "_", text)
    text = text.strip("._")
    return (text or fallback)[:100]


def _json_default(value: object):
    if isinstance(value, datetime):
        return value.isoformat()
    if hasattr(value, "value"):
        return value.value
    return str(value)


def _json_bytes(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, indent=2, default=_json_default).encode("utf-8")


def _dt(value: datetime | None) -> str:
    return value.isoformat() if value else ""


def _enum_value(value: object) -> object:
    return value.value if hasattr(value, "value") else value


def _object_suffix(object_key: str | None, default: str = ".bin") -> str:
    suffix = Path(object_key or "").suffix.lower()
    return suffix or default


def _filename(value: str | None, fallback: str) -> str:
    return _safe_path_part(value or fallback, fallback)


def _read_stored_object(object_key: str | None) -> bytes | None:
    if not object_key:
        return None
    settings = get_settings()
    try:
        client = get_minio_client()
        resp = client.get_object(settings.minio_bucket, object_key)
        try:
            return resp.read()
        finally:
            resp.close()
            resp.release_conn()
    except Exception:
        fallback = settings.uploads_dir / object_key.replace("/", "_")
        if fallback.exists():
            return fallback.read_bytes()
        return None


def _student_dir(class_code: str | None, code: str | None, name: str | None) -> str:
    return "_".join([
        _safe_path_part(class_code, "no_class"),
        _safe_path_part(code, "no_code"),
        _safe_path_part(name, "no_name"),
    ])


def _criterion_score_dict(score: CriterionScore) -> dict:
    return {
        "id": score.id,
        "code": score.code,
        "score": score.score,
        "max_score": score.max_score,
        "not_applicable": score.not_applicable,
        "reasoning": score.reasoning,
        "issues": [
            {
                "id": issue.id,
                "element_refs": issue.element_refs,
                "description": issue.description,
                "severity": issue.severity,
                "suggestion": issue.suggestion,
            }
            for issue in score.issues
        ],
    }


def _grading_result_dict(result: GradingResult) -> dict:
    return {
        "id": result.id,
        "run_id": result.run_id,
        "submission_id": result.submission_id,
        "rubric_version_id": result.rubric_version_id,
        "model_key": result.model_key,
        "status": _enum_value(result.status),
        "total_score": result.total_score,
        "max_total_score": result.max_total_score,
        "normalized_score": result.normalized_score,
        "overall_comment": result.overall_comment,
        "raw_response": result.raw_response,
        "prompt_tokens": result.prompt_tokens,
        "completion_tokens": result.completion_tokens,
        "elapsed_seconds": result.elapsed_seconds,
        "error": result.error,
        "started_at": _dt(result.started_at),
        "finished_at": _dt(result.finished_at),
        "criterion_scores": [_criterion_score_dict(c) for c in result.criterion_scores],
    }


def _teacher_score_dict(score: TeacherScore | ChatTeacherScore) -> dict:
    return {
        "id": score.id,
        "teacher_id": score.teacher_id,
        "code": score.code,
        "score": score.score,
        "max_score": score.max_score,
        "not_applicable": score.not_applicable,
        "comment": score.comment,
        "issues": [
            {
                "id": issue.id,
                "element_refs": issue.element_refs,
                "description": issue.description,
                "severity": issue.severity,
            }
            for issue in score.issues
        ],
    }


def _chat_criterion_score_dict(score: ChatCriterionScore) -> dict:
    return {
        "id": score.id,
        "session_id": score.session_id,
        "message_id": score.message_id,
        "code": score.code,
        "score": score.score,
        "max_score": score.max_score,
        "not_applicable": score.not_applicable,
        "has_issue": score.has_issue,
        "brief": score.brief,
    }


def _teacher_total(scores: list[TeacherScore] | list[ChatTeacherScore]) -> str:
    if not scores:
        return ""
    total = sum(s.score for s in scores if not s.not_applicable)
    max_total = sum(s.max_score for s in scores if not s.not_applicable)
    return f"{total:.2f}/{max_total:.2f}"


def _best_auto_score(results: list[GradingResult]) -> float | None:
    values = [
        r.normalized_score
        for r in results
        if _enum_value(r.status) == "success"
    ]
    return max(values) if values else None


@router.get("/export-uml-artifacts")
async def export_uml_artifacts(
    class_id: int | None = None,
    class_code: str | None = None,
    experiment_group: str | None = None,
    session_group: str | None = None,
    cohort_group: str | None = None,
    diagram_type: str | None = None,
    user_id: int | None = None,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    now = datetime.now(timezone.utc)
    filter_class_code = class_code
    zip_buf = BytesIO()
    manifest_rows: list[dict[str, object]] = []
    missing_files: list[dict[str, object]] = []

    users = (await session.execute(
        select(AppUser).where(AppUser.role == "student")
    )).scalars().all()
    user_by_username = {u.username: u for u in users}

    submissions = (await session.execute(
        select(Submission)
        .options(
            selectinload(Submission.student).selectinload(Student.class_),
            selectinload(Submission.grading_results)
            .selectinload(GradingResult.criterion_scores)
            .selectinload(CriterionScore.issues),
            selectinload(Submission.teacher_scores).selectinload(TeacherScore.issues),
        )
        .order_by(Submission.id)
    )).scalars().all()

    chat_rows = (await session.execute(
        select(ChatSession, AppUser, SchoolClass)
        .join(AppUser, ChatSession.user_id == AppUser.id)
        .outerjoin(SchoolClass, AppUser.class_id == SchoolClass.id)
        .options(
            selectinload(ChatSession.messages).selectinload(ChatMessage.criterion_scores),
            selectinload(ChatSession.teacher_scores).selectinload(ChatTeacherScore.issues),
        )
        .order_by(AppUser.username, ChatSession.id)
    )).all()

    exported_traditional = 0
    exported_chat = 0

    with zipfile.ZipFile(zip_buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for sub in submissions:
            stu = sub.student
            cls = stu.class_ if stu else None
            user = user_by_username.get(stu.student_code) if stu else None
            cls_code = cls.code if cls else ""
            student_code = stu.student_code if stu else ""
            student_name = stu.name if stu else ""
            if class_id and (not cls or cls.id != class_id):
                continue
            if filter_class_code and cls and filter_class_code != cls.code:
                continue
            if filter_class_code and not cls:
                continue
            if experiment_group and (not user or user.experiment_group != experiment_group):
                continue
            if session_group and (not user or user.session_group != session_group):
                continue
            if cohort_group and (not user or user.cohort_group != cohort_group):
                continue
            if diagram_type and sub.diagram_type != diagram_type:
                continue
            if user_id and (not user or user.id != user_id):
                continue
            exported_traditional += 1
            base = (
                "traditional_submissions/"
                f"{_student_dir(cls_code, student_code, student_name)}/"
                f"submission_{sub.id}_{_safe_path_part(sub.diagram_type)}_{_safe_path_part(sub.problem_code, 'no_problem')}"
            )
            mdj_name = _filename(sub.original_filename, f"submission_{sub.id}.mdj")
            if not mdj_name.lower().endswith(".mdj"):
                mdj_name += ".mdj"
            zf.writestr(f"{base}/{mdj_name}", _json_bytes(sub.raw_data))
            zf.writestr(f"{base}/ir.json", _json_bytes(sub.ir_data))
            zf.writestr(
                f"{base}/auto_scores.json",
                _json_bytes([_grading_result_dict(r) for r in sub.grading_results]),
            )
            zf.writestr(
                f"{base}/teacher_scores.json",
                _json_bytes([_teacher_score_dict(r) for r in sub.teacher_scores]),
            )
            zf.writestr(
                f"{base}/metadata.json",
                _json_bytes({
                    "source": "traditional_submission",
                    "submission_id": sub.id,
                    "student_id": sub.student_id,
                    "student_code": student_code,
                    "student_name": student_name,
                    "class_code": cls_code,
                    "experiment_group": user.experiment_group if user else None,
                    "cohort_group": user.cohort_group if user else None,
                    "problem_order": user.problem_order if user else None,
                    "diagram_type": sub.diagram_type,
                    "problem_code": sub.problem_code,
                    "problem_statement": sub.problem_statement,
                    "original_filename": sub.original_filename,
                    "image_object_key": sub.image_filename,
                    "source_format": sub.source_format,
                    "element_count": sub.element_count,
                    "relation_count": sub.relation_count,
                    "container_count": sub.container_count,
                    "parse_error": sub.parse_error,
                    "created_at": _dt(sub.created_at),
                    "updated_at": _dt(sub.updated_at),
                }),
            )

            image_zip_path = ""
            if sub.image_filename:
                image_data = _read_stored_object(sub.image_filename)
                if image_data:
                    image_zip_path = f"{base}/image{_object_suffix(sub.image_filename)}"
                    zf.writestr(image_zip_path, image_data)
                else:
                    missing_files.append({
                        "source": "traditional_submission",
                        "submission_id": sub.id,
                        "object_key": sub.image_filename,
                    })

            manifest_rows.append({
                "source": "traditional_submission",
                "student_code": student_code,
                "student_name": student_name,
                "class_code": cls_code,
                "experiment_group": user.experiment_group if user else "",
                "cohort_group": user.cohort_group if user else "",
                "problem_order": user.problem_order if user else "",
                "diagram_type": sub.diagram_type,
                "problem_code": sub.problem_code,
                "submission_id": sub.id,
                "session_id": "",
                "message_id": "",
                "mdj_path": f"{base}/{mdj_name}",
                "image_path": image_zip_path,
                "auto_score": _best_auto_score(list(sub.grading_results)),
                "teacher_score": _teacher_total(list(sub.teacher_scores)),
                "created_at": _dt(sub.created_at),
                "parse_error": sub.parse_error,
            })

        for cs, user, cls in chat_rows:
            cls_code = cls.code if cls else ""
            student_code = user.username
            student_name = user.display_name
            if class_id and (not cls or cls.id != class_id):
                continue
            if filter_class_code and cls and filter_class_code != cls.code:
                continue
            if filter_class_code and not cls:
                continue
            if experiment_group and user.experiment_group != experiment_group:
                continue
            if session_group and (user.session_group or cs.experiment_group) != session_group:
                continue
            if cohort_group and user.cohort_group != cohort_group:
                continue
            if diagram_type and cs.diagram_type != diagram_type:
                continue
            if user_id and user.id != user_id:
                continue
            exported_chat += 1
            base = (
                "chat_sessions/"
                f"{_student_dir(cls_code, student_code, student_name)}/"
                f"session_{cs.id}_{_safe_path_part(cs.diagram_type)}_{_safe_path_part(cs.problem_code, 'no_problem')}"
            )
            messages_payload = [
                {
                    "id": msg.id,
                    "session_id": msg.session_id,
                    "role": msg.role,
                    "content": msg.content,
                    "mdj_filename": msg.mdj_filename,
                    "mdj_object_key": msg.mdj_object_key,
                    "image_filename": msg.image_filename,
                    "image_object_key": msg.image_object_key,
                    "score_snapshot": msg.score_snapshot,
                    "prompt_tokens": msg.prompt_tokens,
                    "completion_tokens": msg.completion_tokens,
                    "created_at": _dt(msg.created_at),
                    "criterion_scores": [
                        _chat_criterion_score_dict(score)
                        for score in msg.criterion_scores
                    ],
                }
                for msg in cs.messages
            ]
            zf.writestr(f"{base}/messages.json", _json_bytes(messages_payload))
            zf.writestr(
                f"{base}/teacher_scores.json",
                _json_bytes([_teacher_score_dict(r) for r in cs.teacher_scores]),
            )
            zf.writestr(f"{base}/latest_ir.json", _json_bytes(cs.latest_ir_data))
            zf.writestr(
                f"{base}/metadata.json",
                _json_bytes({
                    "source": "chat_session",
                    "session_id": cs.id,
                    "user_id": cs.user_id,
                    "student_code": student_code,
                    "student_name": student_name,
                    "class_code": cls_code,
                    "experiment_group": user.experiment_group,
                    "user_session_group": user.session_group,
                    "session_experiment_group": cs.experiment_group,
                    "cohort_group": user.cohort_group,
                    "problem_order": user.problem_order,
                    "diagram_type": cs.diagram_type,
                    "problem_code": cs.problem_code,
                    "problem_statement": cs.problem_statement,
                    "status": cs.status,
                    "round_number": cs.round_number,
                    "latest_score": cs.latest_score,
                    "passed": cs.passed,
                    "total_active_seconds": cs.total_active_seconds,
                    "created_at": _dt(cs.created_at),
                    "updated_at": _dt(cs.updated_at),
                }),
            )

            for msg in cs.messages:
                mdj_zip_path = ""
                image_zip_path = ""
                if msg.mdj_object_key:
                    mdj_data = _read_stored_object(msg.mdj_object_key)
                    mdj_name = _filename(msg.mdj_filename, f"message_{msg.id}.mdj")
                    if not mdj_name.lower().endswith(".mdj"):
                        mdj_name += ".mdj"
                    if mdj_data:
                        mdj_zip_path = f"{base}/message_{msg.id}_{mdj_name}"
                        zf.writestr(mdj_zip_path, mdj_data)
                    else:
                        missing_files.append({
                            "source": "chat_session",
                            "session_id": cs.id,
                            "message_id": msg.id,
                            "object_key": msg.mdj_object_key,
                        })
                if msg.image_object_key:
                    image_data = _read_stored_object(msg.image_object_key)
                    image_name = _filename(
                        msg.image_filename,
                        f"message_{msg.id}{_object_suffix(msg.image_object_key)}",
                    )
                    if "." not in image_name:
                        image_name += _object_suffix(msg.image_object_key)
                    if image_data:
                        image_zip_path = f"{base}/message_{msg.id}_{image_name}"
                        zf.writestr(image_zip_path, image_data)
                    else:
                        missing_files.append({
                            "source": "chat_session",
                            "session_id": cs.id,
                            "message_id": msg.id,
                            "object_key": msg.image_object_key,
                        })

                if msg.role == "user" and (msg.mdj_object_key or msg.image_object_key):
                    manifest_rows.append({
                        "source": "chat_session",
                        "student_code": student_code,
                        "student_name": student_name,
                        "class_code": cls_code,
                        "experiment_group": user.experiment_group or "",
                        "session_group": user.session_group or cs.experiment_group or "",
                        "cohort_group": user.cohort_group or "",
                        "problem_order": user.problem_order or "",
                        "diagram_type": cs.diagram_type,
                        "problem_code": cs.problem_code,
                        "submission_id": "",
                        "session_id": cs.id,
                        "message_id": msg.id,
                        "mdj_path": mdj_zip_path,
                        "image_path": image_zip_path,
                        "auto_score": msg.score_snapshot if msg.score_snapshot is not None else cs.latest_score,
                        "teacher_score": _teacher_total(list(cs.teacher_scores)),
                        "created_at": _dt(msg.created_at),
                        "parse_error": "",
                    })

        csv_buf = StringIO()
        fieldnames = [
            "source", "student_code", "student_name", "class_code",
            "experiment_group", "session_group", "cohort_group", "problem_order",
            "diagram_type", "problem_code", "submission_id", "session_id",
            "message_id", "mdj_path", "image_path", "auto_score",
            "teacher_score", "created_at", "parse_error",
        ]
        writer = csv.DictWriter(csv_buf, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(manifest_rows)
        zf.writestr("manifest.csv", "\ufeff" + csv_buf.getvalue())
        zf.writestr("manifest.json", _json_bytes(manifest_rows))
        zf.writestr("missing_files.json", _json_bytes(missing_files))
        zf.writestr(
            "summary.json",
            _json_bytes({
                "exported_at": now,
                "filters": {
                    "class_id": class_id,
                    "class_code": filter_class_code,
                    "experiment_group": experiment_group,
                    "session_group": session_group,
                    "cohort_group": cohort_group,
                    "diagram_type": diagram_type,
                    "user_id": user_id,
                },
                "traditional_submissions": exported_traditional,
                "chat_sessions": exported_chat,
                "uml_items": len(manifest_rows),
                "missing_files": len(missing_files),
            }),
        )

    zip_buf.seek(0)
    filename = f"uml_artifacts_{now.strftime('%Y%m%d_%H%M%S')}.zip"
    quoted = quote(filename)
    return StreamingResponse(
        zip_buf,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"},
    )


class ImportRow(BaseModel):
    username: str
    display_name: str
    class_code: str
    status: str  # created / exists / error
    message: str = ""


class ImportResult(BaseModel):
    total: int
    created: int
    skipped: int
    errors: int
    rows: list[ImportRow]


@router.post("/import-students", response_model=ImportResult)
async def import_students(
    file: UploadFile = File(..., description="xlsx 文件，列：班级、学号、姓名"),
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """从 xlsx 批量导入学生账号。

    xlsx 格式（第一行为表头）：
    | 班级 | 学号 | 姓名 | 分组(可选) | C组编号(可选) | 题目顺序(可选) |

    - 用户名 = 学号
    - 默认密码 = 学号 + "123"
    - 角色 = student
    - 班级不存在时自动创建
    - 第 4 列可选：experiment_group（agent_a / agent_b / agent_c / control）
    - 第 5 列可选：cohort_group（C1 ~ C16）
    - 第 6 列可选：problem_order（normal / reversed）
    """
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "请上传 .xlsx 文件")

    raw = await file.read()
    try:
        wb = load_workbook(BytesIO(raw), read_only=True)
    except Exception as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"无法解析 xlsx: {e}")

    ws = wb.active
    if ws is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "xlsx 中无工作表")

    rows_out: list[ImportRow] = []
    created = 0
    skipped = 0
    errors = 0

    # 缓存班级 code -> id
    class_cache: dict[str, int] = {}

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or len(row) < 3:
            continue

        class_code = str(row[0] or "").strip()
        student_code = str(row[1] or "").strip()
        name = str(row[2] or "").strip()
        exp_group = str(row[3] or "").strip().lower() if len(row) > 3 and row[3] else None
        if exp_group and exp_group not in ("agent_a", "agent_b", "agent_c", "control"):
            exp_group = None  # 忽略无效分组
        cohort = str(row[4] or "").strip().upper() if len(row) > 4 and row[4] else None
        if cohort and not (cohort.startswith("C") and cohort[1:].isdigit()):
            cohort = None
        prob_order = str(row[5] or "").strip().lower() if len(row) > 5 and row[5] else None
        if prob_order and prob_order not in ("normal", "reversed"):
            prob_order = None

        if not class_code or not student_code:
            rows_out.append(ImportRow(
                username=student_code, display_name=name,
                class_code=class_code, status="error",
                message=f"第 {i} 行：班级或学号为空",
            ))
            errors += 1
            continue

        # 确保班级存在
        if class_code not in class_cache:
            cls_ = (await session.execute(
                select(SchoolClass).where(SchoolClass.code == class_code)
            )).scalar_one_or_none()
            if cls_ is None:
                cls_ = SchoolClass(code=class_code, name=class_code)
                session.add(cls_)
                await session.flush()
            class_cache[class_code] = cls_.id

        class_id = class_cache[class_code]

        # 检查用户是否已存在
        existing = (await session.execute(
            select(AppUser).where(AppUser.username == student_code)
        )).scalar_one_or_none()

        if existing:
            rows_out.append(ImportRow(
                username=student_code, display_name=name,
                class_code=class_code, status="exists",
                message="账号已存在，已跳过",
            ))
            skipped += 1
            continue

        user = AppUser(
            username=student_code,
            display_name=name or student_code,
            hashed_password=hash_password(student_code + "123"),
            role="student",
            is_active=True,
            class_id=class_id,
            experiment_group=exp_group,
            cohort_group=cohort,
            problem_order=prob_order,
        )
        session.add(user)
        rows_out.append(ImportRow(
            username=student_code, display_name=name,
            class_code=class_code, status="created",
        ))
        created += 1

    await session.commit()

    return ImportResult(
        total=len(rows_out),
        created=created,
        skipped=skipped,
        errors=errors,
        rows=rows_out,
    )


class AssignGroupIn(BaseModel):
    """按班级分配实验分组。"""
    class_code: str
    experiment_group: str  # agent_a | agent_b | agent_c | control


class AssignGroupResult(BaseModel):
    class_code: str
    experiment_group: str
    updated: int


VALID_GROUPS = {"agent_a", "agent_b", "agent_c", "control", "multi_model_a"}


@router.post("/assign-group", response_model=AssignGroupResult)
async def assign_experiment_group(
    body: AssignGroupIn,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """按班级批量分配实验分组。

    将指定班级的所有学生设为相同的 experiment_group。
    """
    if body.experiment_group not in VALID_GROUPS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"无效分组，可选: {', '.join(sorted(VALID_GROUPS))}",
        )

    cls_ = (await session.execute(
        select(SchoolClass).where(SchoolClass.code == body.class_code)
    )).scalar_one_or_none()
    if cls_ is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"班级 {body.class_code} 不存在")

    # 查找该班级的所有学生用户
    users = (await session.execute(
        select(AppUser).where(
            AppUser.class_id == cls_.id,
            AppUser.role == "student",
        )
    )).scalars().all()

    for u in users:
        u.experiment_group = body.experiment_group
        session.add(u)

    await session.commit()

    return AssignGroupResult(
        class_code=body.class_code,
        experiment_group=body.experiment_group,
        updated=len(users),
    )


class AssignSessionGroupIn(BaseModel):
    """批量分配当次实验 Agent 分组（session_group）。

    scope:
      - "all"          所有学生
      - "class"        指定 class_code 班级
      - "users"        指定 user_ids 列表
    """
    scope: str = "all"
    class_code: str | None = None
    user_ids: list[int] | None = None
    session_group: str  # agent_a | agent_b | agent_c | control | ""(清空)


class AssignSessionGroupResult(BaseModel):
    scope: str
    session_group: str | None
    updated: int


@router.post("/assign-session-group", response_model=AssignSessionGroupResult)
async def assign_session_group(
    body: AssignSessionGroupIn,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """批量设置学生的 session_group（当次实验使用的 Agent 类型）。

    与历史 experiment_group 解耦，可在每次实验前重新分配。
    传 session_group="" 表示清空。
    """
    target: str | None
    if body.session_group == "":
        target = None
    elif body.session_group in VALID_GROUPS:
        target = body.session_group
    else:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"无效分组，可选: {', '.join(sorted(VALID_GROUPS))} 或空字符串",
        )

    q = select(AppUser).where(AppUser.role == "student")
    if body.scope == "class":
        if not body.class_code:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "scope=class 时必须提供 class_code")
        cls_ = (await session.execute(
            select(SchoolClass).where(SchoolClass.code == body.class_code)
        )).scalar_one_or_none()
        if cls_ is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, f"班级不存在: {body.class_code}")
        q = q.where(AppUser.class_id == cls_.id)
    elif body.scope == "users":
        if not body.user_ids:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "scope=users 时必须提供 user_ids")
        q = q.where(AppUser.id.in_(body.user_ids))
    elif body.scope != "all":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"无效 scope: {body.scope}")

    users = (await session.execute(q)).scalars().all()
    for u in users:
        u.session_group = target
        session.add(u)
    await session.commit()

    return AssignSessionGroupResult(
        scope=body.scope,
        session_group=target,
        updated=len(users),
    )


@router.get("/groups-summary")
async def groups_summary(
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """查看各实验分组的学生数量。"""
    from sqlalchemy import func

    rows = (await session.execute(
        select(
            AppUser.experiment_group,
            func.count(AppUser.id),
        ).where(
            AppUser.role == "student",
        ).group_by(AppUser.experiment_group)
    )).all()

    return {
        "groups": [
            {"group": g or "unassigned", "count": c}
            for g, c in rows
        ]
    }


@router.get("/groups-members")
async def groups_members(
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """查看各实验分组的成员列表，含聊天统计。"""
    # 查所有学生及其班级
    users = (await session.execute(
        select(AppUser, SchoolClass.code, SchoolClass.name)
        .outerjoin(SchoolClass, AppUser.class_id == SchoolClass.id)
        .where(AppUser.role == "student")
        .order_by(AppUser.experiment_group, AppUser.username)
    )).all()

    # 查每个学生的聊天会话统计
    chat_stats_rows = (await session.execute(
        select(
            ChatSession.user_id,
            func.count(ChatSession.id),
            func.max(ChatSession.latest_score),
            func.max(ChatSession.updated_at),
        ).group_by(ChatSession.user_id)
    )).all()
    chat_stats = {
        uid: {"session_count": cnt, "best_score": score, "last_active": ts}
        for uid, cnt, score, ts in chat_stats_rows
    }

    result: dict[str, list] = {}
    for u, cls_code, cls_name in users:
        g = u.experiment_group or "unassigned"
        cs = chat_stats.get(u.id, {})
        result.setdefault(g, []).append({
            "id": u.id,
            "username": u.username,
            "display_name": u.display_name,
            "class_code": cls_code or "",
            "class_name": cls_name or "",
            "session_count": cs.get("session_count", 0),
            "best_score": cs.get("best_score"),
            "last_active": cs.get("last_active"),
        })

    return {"groups": result}


# ---------- 管理员查看聊天会话 ----------

@router.get("/chat-sessions")
async def admin_list_chat_sessions(
    experiment_group: str | None = None,
    session_group: str | None = None,
    class_id: int | None = None,
    class_code: str | None = None,
    diagram_type: str | None = None,
    user_id: int | None = None,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """管理员查看所有聊天会话，可按分组 / 用户筛选。"""
    q = (
        select(
            ChatSession,
            AppUser.username,
            AppUser.display_name,
            AppUser.experiment_group,
            AppUser.session_group,
            SchoolClass.code,
        )
        .join(AppUser, ChatSession.user_id == AppUser.id)
        .outerjoin(SchoolClass, AppUser.class_id == SchoolClass.id)
        .order_by(ChatSession.updated_at.desc())
    )
    if experiment_group:
        q = q.where(AppUser.experiment_group == experiment_group)
    if session_group:
        q = q.where((AppUser.session_group == session_group) | (ChatSession.experiment_group == session_group))
    if class_id:
        q = q.where(AppUser.class_id == class_id)
    if class_code:
        q = q.where(SchoolClass.code == class_code)
    if diagram_type:
        q = q.where(ChatSession.diagram_type == diagram_type)
    if user_id:
        q = q.where(ChatSession.user_id == user_id)

    rows = (await session.execute(q)).all()

    # 批量查询每个 session 的教师评分数量
    session_ids = [cs.id for cs, *_ in rows]
    scored_counts: dict[int, int] = {}
    if session_ids:
        cnt_q = (
            select(
                ChatTeacherScore.session_id,
                func.count(ChatTeacherScore.id),
            )
            .where(ChatTeacherScore.session_id.in_(session_ids))
            .group_by(ChatTeacherScore.session_id)
        )
        for sid, cnt in (await session.execute(cnt_q)).all():
            scored_counts[sid] = cnt

    return [
        {
            "id": cs.id,
            "user_id": cs.user_id,
            "username": uname,
            "display_name": dname,
            "experiment_group": eg or "",
            "session_group": sg or cs.experiment_group or "",
            "class_code": cls_code or "",
            "diagram_type": cs.diagram_type,
            "problem_code": cs.problem_code,
            "problem_statement": cs.problem_statement,
            "status": cs.status,
            "round_number": cs.round_number,
            "latest_score": cs.latest_score,
            "passed": cs.passed,
            "teacher_scored_count": scored_counts.get(cs.id, 0),
            "created_at": cs.created_at,
            "updated_at": cs.updated_at,
        }
        for cs, uname, dname, eg, sg, cls_code in rows
    ]


@router.get("/chat-sessions/{session_id}")
async def admin_get_chat_session(
    session_id: int,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """管理员查看单个聊天会话详情（含消息）。"""
    cs = await session.get(
        ChatSession, session_id,
        options=[selectinload(ChatSession.messages)],
    )
    if cs is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "会话不存在")

    user = await session.get(AppUser, cs.user_id)

    messages_out = []
    for m in cs.messages:
        img_url = None
        if m.image_object_key:
            try:
                img_url = get_minio_url(m.image_object_key)
            except Exception:
                pass
        messages_out.append({
            "id": m.id,
            "session_id": m.session_id,
            "role": m.role,
            "content": m.content,
            "mdj_filename": m.mdj_filename,
            "image_filename": m.image_filename,
            "image_url": img_url,
            "score_snapshot": m.score_snapshot,
            "created_at": m.created_at,
        })

    return {
        "id": cs.id,
        "user_id": cs.user_id,
        "username": user.username if user else "",
        "display_name": user.display_name if user else "",
        "experiment_group": cs.experiment_group,
        "diagram_type": cs.diagram_type,
        "problem_code": cs.problem_code,
        "problem_statement": cs.problem_statement,
        "status": cs.status,
        "round_number": cs.round_number,
        "latest_score": cs.latest_score,
        "passed": cs.passed,
        "created_at": cs.created_at,
        "updated_at": cs.updated_at,
        "messages": messages_out,
    }


# ---------- 教师对聊天会话逐条指标评分 ----------


@router.post("/chat-sessions/{session_id}/scores", response_model=ChatTeacherScoreOut)
async def upsert_chat_teacher_score(
    session_id: int,
    body: ChatTeacherScoreIn,
    session: AsyncSession = Depends(get_session),
    user=Depends(require_admin),
):
    """教师对某个聊天会话的某条指标 upsert 评分。"""
    cs = await session.get(ChatSession, session_id)
    if cs is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "会话不存在")
    if body.session_id != session_id:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "session_id 不匹配")

    stmt = select(ChatTeacherScore).where(
        ChatTeacherScore.session_id == session_id,
        ChatTeacherScore.teacher_id == user.id,
        ChatTeacherScore.code == body.code,
    ).options(selectinload(ChatTeacherScore.issues))
    ts = (await session.execute(stmt)).scalar_one_or_none()

    if ts is None:
        ts = ChatTeacherScore(
            session_id=session_id,
            teacher_id=user.id,
            code=body.code,
        )
        session.add(ts)

    ts.score = body.score
    ts.max_score = body.max_score
    ts.not_applicable = body.not_applicable
    ts.comment = body.comment

    # 清空旧 issue 重新写
    if ts.id:
        for old in list(ts.issues):
            await session.delete(old)

    await session.flush()
    for i in body.issues:
        session.add(ChatTeacherIssue(
            chat_teacher_score_id=ts.id,
            element_refs=list(i.element_refs),
            description=i.description,
            severity=i.severity,
        ))
    await session.commit()

    # 重新载入
    ts = (await session.execute(
        select(ChatTeacherScore).where(ChatTeacherScore.id == ts.id)
        .options(selectinload(ChatTeacherScore.issues))
    )).scalar_one()
    return {
        **ChatTeacherScoreOut.model_validate(ts).model_dump(exclude={"issues"}),
        "issues": [ChatTeacherIssueOut.model_validate(i).model_dump() for i in ts.issues],
    }


@router.get("/chat-sessions/{session_id}/scores", response_model=list[ChatTeacherScoreOut])
async def list_chat_teacher_scores(
    session_id: int,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """获取某聊天会话的全部教师评分。"""
    stmt = (
        select(ChatTeacherScore)
        .where(ChatTeacherScore.session_id == session_id)
        .options(selectinload(ChatTeacherScore.issues))
        .order_by(ChatTeacherScore.code)
    )
    rows = (await session.execute(stmt)).scalars().all()
    return [
        {
            **ChatTeacherScoreOut.model_validate(r).model_dump(exclude={"issues"}),
            "issues": [ChatTeacherIssueOut.model_validate(i).model_dump() for i in r.issues],
        }
        for r in rows
    ]


# ---------- 学生管理 ----------

class StudentOut(BaseModel):
    id: int
    username: str
    display_name: str
    class_code: str | None = None
    class_name: str | None = None
    experiment_group: str | None = None
    cohort_group: str | None = None
    problem_order: str | None = None
    session_group: str | None = None
    is_active: bool = True
    session_count: int = 0
    best_score: float | None = None

    class Config:
        from_attributes = True


class StudentUpdateIn(BaseModel):
    display_name: str | None = None
    experiment_group: str | None = None
    cohort_group: str | None = None
    problem_order: str | None = None
    session_group: str | None = None
    is_active: bool | None = None


@router.get("/students", response_model=list[StudentOut])
async def list_students(
    search: str | None = None,
    class_code: str | None = None,
    experiment_group: str | None = None,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """管理员查看学生列表，支持搜索和筛选。"""
    q = (
        select(
            AppUser,
            SchoolClass.code.label("cls_code"),
            SchoolClass.name.label("cls_name"),
        )
        .outerjoin(SchoolClass, AppUser.class_id == SchoolClass.id)
        .where(AppUser.role == "student")
        .order_by(SchoolClass.code, AppUser.username)
    )
    if search:
        like = f"%{search}%"
        q = q.where(
            AppUser.username.ilike(like) | AppUser.display_name.ilike(like)
        )
    if class_code:
        q = q.where(SchoolClass.code == class_code)
    if experiment_group:
        if experiment_group == "unassigned":
            q = q.where(AppUser.experiment_group.is_(None))
        else:
            q = q.where(AppUser.experiment_group == experiment_group)

    rows = (await session.execute(q)).all()

    # 聊天统计
    chat_rows = (await session.execute(
        select(
            ChatSession.user_id,
            func.count(ChatSession.id),
            func.max(ChatSession.latest_score),
        ).group_by(ChatSession.user_id)
    )).all()
    chat_map = {uid: (cnt, sc) for uid, cnt, sc in chat_rows}

    return [
        StudentOut(
            id=u.id,
            username=u.username,
            display_name=u.display_name,
            class_code=cc,
            class_name=cn,
            experiment_group=u.experiment_group,
            cohort_group=u.cohort_group,
            problem_order=u.problem_order,
            session_group=u.session_group,
            is_active=u.is_active,
            session_count=chat_map.get(u.id, (0, None))[0],
            best_score=chat_map.get(u.id, (0, None))[1],
        )
        for u, cc, cn in rows
    ]


@router.put("/students/{user_id}", response_model=StudentOut)
async def update_student(
    user_id: int,
    body: StudentUpdateIn,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """修改学生信息（分组、C组、题序、姓名、激活状态）。"""
    user = await session.get(AppUser, user_id)
    if user is None or user.role != "student":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "学生不存在")

    if body.display_name is not None:
        user.display_name = body.display_name
    if body.experiment_group is not None:
        if body.experiment_group == "":
            user.experiment_group = None
        elif body.experiment_group not in VALID_GROUPS:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, f"无效分组: {body.experiment_group}")
        else:
            user.experiment_group = body.experiment_group
    if body.cohort_group is not None:
        user.cohort_group = body.cohort_group if body.cohort_group else None
    if body.problem_order is not None:
        if body.problem_order and body.problem_order not in ("normal", "reversed", ""):
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "题序只能是 normal 或 reversed")
        user.problem_order = body.problem_order if body.problem_order else None
    if body.session_group is not None:
        if body.session_group == "":
            user.session_group = None
        elif body.session_group not in VALID_GROUPS:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, f"无效分组: {body.session_group}")
        else:
            user.session_group = body.session_group
    if body.is_active is not None:
        user.is_active = body.is_active

    session.add(user)
    await session.commit()
    await session.refresh(user)

    # 获取班级信息
    cls_ = await session.get(SchoolClass, user.class_id) if user.class_id else None

    return StudentOut(
        id=user.id,
        username=user.username,
        display_name=user.display_name,
        class_code=cls_.code if cls_ else None,
        class_name=cls_.name if cls_ else None,
        experiment_group=user.experiment_group,
        cohort_group=user.cohort_group,
        problem_order=user.problem_order,
        session_group=user.session_group,
        is_active=user.is_active,
    )


@router.post("/students/{user_id}/reset-password")
async def reset_student_password(
    user_id: int,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """重置学生密码为 学号+123。"""
    user = await session.get(AppUser, user_id)
    if user is None or user.role != "student":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "学生不存在")

    user.hashed_password = hash_password(user.username + "123")
    session.add(user)
    await session.commit()

    return {"ok": True, "message": f"已重置 {user.username} 的密码为 {user.username}123"}


@router.delete("/students/{user_id}")
async def delete_student(
    user_id: int,
    session: AsyncSession = Depends(get_session),
    _u=Depends(require_admin),
):
    """删除学生账号（关联的聊天记录会级联删除）。"""
    user = await session.get(AppUser, user_id)
    if user is None or user.role != "student":
        raise HTTPException(status.HTTP_404_NOT_FOUND, "学生不存在")

    await session.delete(user)
    await session.commit()

    return {"ok": True, "message": f"已删除 {user.username}"}
