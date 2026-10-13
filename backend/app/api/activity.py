"""学生使用时长追踪 API（心跳上报 + 时长查询）。"""

from __future__ import annotations

from io import BytesIO
from datetime import datetime, timedelta, timezone
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import func, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from ..core.db import get_session
from ..core.security import get_current_user, require_role
from ..models.activity import UserActivityLog
from ..models.chat import ChatSession
from ..models.user import AppUser

router = APIRouter()

# ---- 配置 ----
HEARTBEAT_INTERVAL = 30  # 前端每 30 秒发一次
MERGE_GAP = 90  # 同一块内两次心跳最大间隔秒数


# ---- Schema ----
class HeartbeatIn(BaseModel):
    page: str = ""  # 当前页面路径
    chat_session_id: int | None = None
    diagram_type: str | None = None


class HeartbeatOut(BaseModel):
    ok: bool = True


class MyActivitySummary(BaseModel):
    total_seconds: int
    today_seconds: int
    by_session: list[dict[str, Any]]


class AdminActivityItem(BaseModel):
    user_id: int
    username: str
    display_name: str
    class_code: str | None
    experiment_group: str | None
    total_seconds: int
    heartbeat_seconds: int = 0
    reconstructed_seconds: int = 0
    session_count: int
    by_session: list[dict[str, Any]] | None = None


# ---- 心跳端点（学生） ----
@router.post("/heartbeat", response_model=HeartbeatOut)
async def heartbeat(
    body: HeartbeatIn,
    session: AsyncSession = Depends(get_session),
    current_user: AppUser = Depends(get_current_user),
):
    """
    前端每 30s 调用一次。服务端将相邻心跳合并为 UserActivityLog 条目：
    - 如果存在一个 ended_at 在 MERGE_GAP 秒内的记录 → 延展该条目
    - 否则 → 新建条目
    """
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(seconds=MERGE_GAP)

    # 查找可以合并的最近一条活动记录
    stmt = (
        select(UserActivityLog)
        .where(
            UserActivityLog.user_id == current_user.id,
            UserActivityLog.ended_at >= cutoff,
            UserActivityLog.source == "heartbeat",
        )
        .order_by(UserActivityLog.ended_at.desc())
        .limit(1)
    )
    last_log = (await session.execute(stmt)).scalar_one_or_none()

    if last_log is not None:
        # 延展现有块
        last_log.ended_at = now
        last_log.duration_seconds = int((now - last_log.started_at).total_seconds())
        # 更新上下文
        last_log.context = {
            "page": body.page,
            "diagram_type": body.diagram_type,
            "chat_session_id": body.chat_session_id,
        }
        # 如果有关联 chat_session 也更新
        if body.chat_session_id and last_log.chat_session_id != body.chat_session_id:
            last_log.chat_session_id = body.chat_session_id
        session.add(last_log)
    else:
        # 新建使用块
        new_log = UserActivityLog(
            user_id=current_user.id,
            started_at=now,
            ended_at=now,
            duration_seconds=0,
            chat_session_id=body.chat_session_id,
            context={
                "page": body.page,
                "diagram_type": body.diagram_type,
                "chat_session_id": body.chat_session_id,
            },
            source="heartbeat",
        )
        session.add(new_log)

    # 同步更新对应 ChatSession 的 total_active_seconds
    if body.chat_session_id:
        await _update_session_active_seconds(session, body.chat_session_id, current_user.id)

    await session.commit()
    return HeartbeatOut(ok=True)


async def _update_session_active_seconds(
    session: AsyncSession, chat_session_id: int, user_id: int
):
    """重新计算某个 ChatSession 的总活跃秒数。"""
    stmt = (
        select(func.coalesce(func.sum(UserActivityLog.duration_seconds), 0))
        .where(
            UserActivityLog.user_id == user_id,
            UserActivityLog.chat_session_id == chat_session_id,
        )
    )
    total = (await session.execute(stmt)).scalar_one()
    await session.execute(
        update(ChatSession)
        .where(ChatSession.id == chat_session_id)
        .values(total_active_seconds=total)
    )


# ---- 学生查看自己的使用时长 ----
@router.get("/me", response_model=MyActivitySummary)
async def my_activity(
    session: AsyncSession = Depends(get_session),
    current_user: AppUser = Depends(get_current_user),
):
    """返回当前学生的总使用时长和按会话分布。"""
    # 总时长
    total_stmt = (
        select(func.coalesce(func.sum(UserActivityLog.duration_seconds), 0))
        .where(UserActivityLog.user_id == current_user.id)
    )
    total_seconds = (await session.execute(total_stmt)).scalar_one()

    # 今日时长
    today_start = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
    today_stmt = (
        select(func.coalesce(func.sum(UserActivityLog.duration_seconds), 0))
        .where(
            UserActivityLog.user_id == current_user.id,
            UserActivityLog.started_at >= today_start,
        )
    )
    today_seconds = (await session.execute(today_stmt)).scalar_one()

    # 按 ChatSession 聚合
    by_session_stmt = (
        select(
            ChatSession.id,
            ChatSession.diagram_type,
            ChatSession.problem_code,
            ChatSession.total_active_seconds,
            ChatSession.status,
            ChatSession.latest_score,
        )
        .where(ChatSession.user_id == current_user.id)
        .order_by(ChatSession.id.desc())
    )
    rows = (await session.execute(by_session_stmt)).all()
    by_session = [
        {
            "session_id": r.id,
            "diagram_type": r.diagram_type,
            "problem_code": r.problem_code,
            "active_seconds": r.total_active_seconds,
            "status": r.status,
            "latest_score": r.latest_score,
        }
        for r in rows
    ]

    return MyActivitySummary(
        total_seconds=total_seconds,
        today_seconds=today_seconds,
        by_session=by_session,
    )


# ---- 管理员查看所有学生使用时长 ----
require_non_student = require_role("admin", "teacher", "researcher")


def _format_duration(seconds: int) -> str:
    seconds = max(0, int(seconds or 0))
    minutes = seconds // 60
    hours = minutes // 60
    mins = minutes % 60
    if hours <= 0:
        return f"{mins}分钟"
    return f"{hours}小时{mins}分钟"


def _group_label(group: str | None) -> str:
    labels = {
        "agent_a": "Agent A",
        "agent_b": "Agent B",
        "agent_c": "Agent C",
        "control": "对照组",
        "multi_model_a": "四模型A",
    }
    return labels.get(group or "", group or "未分配")


def _safe_filename_part(value: str) -> str:
    return "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in value).strip("_")


async def _admin_activity_rows(
    session: AsyncSession = Depends(get_session),
    class_code: str | None = None,
    search: str | None = None,
) -> list[AdminActivityItem]:
    from ..models.class_ import SchoolClass

    class_code = (class_code or "").strip()
    search = (search or "").strip()

    stmt = (
        select(
            AppUser.id,
            AppUser.username,
            AppUser.display_name,
            SchoolClass.code.label("class_code"),
            AppUser.experiment_group,
            func.coalesce(func.sum(UserActivityLog.duration_seconds), 0).label("total_seconds"),
            func.coalesce(
                func.sum(UserActivityLog.duration_seconds).filter(
                    UserActivityLog.source == "heartbeat"
                ),
                0,
            ).label("heartbeat_seconds"),
            func.coalesce(
                func.sum(UserActivityLog.duration_seconds).filter(
                    UserActivityLog.source == "reconstructed"
                ),
                0,
            ).label("reconstructed_seconds"),
            func.count(func.distinct(UserActivityLog.chat_session_id)).label("session_count"),
        )
        .outerjoin(UserActivityLog, UserActivityLog.user_id == AppUser.id)
        .outerjoin(SchoolClass, SchoolClass.id == AppUser.class_id)
        .where(AppUser.role == "student")
        .group_by(AppUser.id, AppUser.username, AppUser.display_name, SchoolClass.code, AppUser.experiment_group)
        .order_by(AppUser.username)
    )

    if class_code == "__unassigned__":
        stmt = stmt.where(AppUser.class_id.is_(None))
    elif class_code:
        stmt = stmt.where(SchoolClass.code == class_code)

    if search:
        like = f"%{search}%"
        stmt = stmt.where(AppUser.username.ilike(like) | AppUser.display_name.ilike(like))

    rows = (await session.execute(stmt)).all()
    return [
        AdminActivityItem(
            user_id=r.id,
            username=r.username,
            display_name=r.display_name,
            class_code=r.class_code,
            experiment_group=r.experiment_group,
            total_seconds=r.total_seconds,
            heartbeat_seconds=r.heartbeat_seconds,
            reconstructed_seconds=r.reconstructed_seconds,
            session_count=r.session_count,
        )
        for r in rows
    ]


@router.get("/admin", response_model=list[AdminActivityItem])
async def admin_activity(
    class_code: str | None = Query(default=None),
    search: str | None = Query(default=None),
    session: AsyncSession = Depends(get_session),
    _: AppUser = Depends(require_non_student),
):
    """管理员看板：按学生汇总使用时长，可按班级/关键词筛选。"""
    return await _admin_activity_rows(session=session, class_code=class_code, search=search)


@router.get("/admin/export")
async def export_admin_activity(
    class_code: str | None = Query(default=None),
    search: str | None = Query(default=None),
    session: AsyncSession = Depends(get_session),
    _: AppUser = Depends(require_non_student),
):
    """导出当前筛选条件下的学生使用时长 xlsx。"""
    rows = await _admin_activity_rows(session=session, class_code=class_code, search=search)
    display_class = "未分班" if class_code == "__unassigned__" else (class_code or "全部班级")

    wb = Workbook()
    ws = wb.active
    ws.title = "使用时长"

    headers = [
        "班级",
        "学号",
        "姓名",
        "实验分组",
        "会话数",
        "使用时长",
        "使用时长(小时)",
        "使用时长(分钟)",
        "使用时长(秒)",
    ]
    ws.append(headers)

    for item in rows:
        total_seconds = int(item.total_seconds or 0)
        ws.append([
            item.class_code or "未分班",
            item.username,
            item.display_name,
            _group_label(item.experiment_group),
            item.session_count,
            _format_duration(total_seconds),
            round(total_seconds / 3600, 2),
            round(total_seconds / 60, 2),
            total_seconds,
        ])

    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="0F172A")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_len + 4, 24)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    suffix = _safe_filename_part(display_class) or "all"
    filename = f"使用时长_{suffix}.xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
        },
    )


# ---- 管理员查看某个学生的详细时长 ----
@router.get("/admin/{user_id}", response_model=AdminActivityItem)
async def admin_user_activity(
    user_id: int,
    session: AsyncSession = Depends(get_session),
    _: AppUser = Depends(require_non_student),
):
    """管理员查看某学生的详细使用时长（含按会话分布）。"""
    from ..models.class_ import SchoolClass

    user = (await session.execute(
        select(AppUser).where(AppUser.id == user_id)
    )).scalar_one_or_none()
    if not user:
        from fastapi import HTTPException
        raise HTTPException(404, "User not found")

    cls = None
    if user.class_id:
        cls = (await session.execute(
            select(SchoolClass).where(SchoolClass.id == user.class_id)
        )).scalar_one_or_none()

    # 总时长及来源分项
    source_totals = (await session.execute(
        select(
            func.coalesce(func.sum(UserActivityLog.duration_seconds), 0).label("total_seconds"),
            func.coalesce(
                func.sum(UserActivityLog.duration_seconds).filter(
                    UserActivityLog.source == "heartbeat"
                ),
                0,
            ).label("heartbeat_seconds"),
            func.coalesce(
                func.sum(UserActivityLog.duration_seconds).filter(
                    UserActivityLog.source == "reconstructed"
                ),
                0,
            ).label("reconstructed_seconds"),
        )
        .where(UserActivityLog.user_id == user_id)
    )).one()

    # 按 session 聚合
    by_session_stmt = (
        select(
            ChatSession.id,
            ChatSession.diagram_type,
            ChatSession.problem_code,
            ChatSession.total_active_seconds,
            ChatSession.status,
            ChatSession.latest_score,
        )
        .where(ChatSession.user_id == user_id)
        .order_by(ChatSession.id.desc())
    )
    rows = (await session.execute(by_session_stmt)).all()
    by_session = [
        {
            "session_id": r.id,
            "diagram_type": r.diagram_type,
            "problem_code": r.problem_code,
            "active_seconds": r.total_active_seconds,
            "status": r.status,
            "latest_score": r.latest_score,
        }
        for r in rows
    ]

    session_count = len([r for r in rows if r.total_active_seconds > 0])

    return AdminActivityItem(
        user_id=user.id,
        username=user.username,
        display_name=user.display_name,
        class_code=cls.code if cls else None,
        experiment_group=user.experiment_group,
        total_seconds=source_totals.total_seconds,
        heartbeat_seconds=source_totals.heartbeat_seconds,
        reconstructed_seconds=source_totals.reconstructed_seconds,
        session_count=session_count,
        by_session=by_session,
    )
